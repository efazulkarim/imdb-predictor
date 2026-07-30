"""Build a two-column IEEE-style conference manuscript from the verified repo artifacts.

The supplied A4 template is an ISO/IEC 29500 Strict OOXML document.  python-docx
handles the transitional OOXML dialect, so the builder makes an in-memory-compatible
copy of the template first, then retains its named styles and page geometry.
"""

from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "conference-template-a4.docx"
OUTPUT = Path(__file__).resolve().with_name("imdb_screenplay_ieee_conference_paper.docx")
FIGURES = ROOT / "figures"


def strict_template_to_transitional(source: Path, target: Path) -> None:
    """Convert the template package namespace declarations without changing its styles."""
    strict_prefix = b"http://purl.oclc.org/ooxml/"
    transitional_prefix = b"http://schemas.openxmlformats.org/"
    content_type_replacements = {
        b"application/vnd.ms-word.document.main+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
        ),
        b"application/vnd.ms-word.styles+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"
        ),
        b"application/vnd.ms-word.fontTable+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml"
        ),
        b"application/vnd.ms-word.settings+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"
        ),
        b"application/vnd.ms-word.numbering+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml"
        ),
        b"application/vnd.ms-word.footnotes+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
        ),
        b"application/vnd.ms-word.endnotes+xml": (
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"
        ),
    }
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst:
        for member in src.infolist():
            data = src.read(member.filename)
            if member.filename.endswith((".xml", ".rels")):
                data = data.replace(strict_prefix, transitional_prefix)
                for old, new in content_type_replacements.items():
                    data = data.replace(old, new)
            dst.writestr(member, data)


def style(doc: Document, *names: str):
    """Find a template style by display name or persistent style id."""
    wanted = {name.casefold() for name in names}
    for candidate in doc.styles:
        if candidate.name.casefold() in wanted or candidate.style_id.casefold() in wanted:
            return candidate
    raise KeyError(f"Template style missing: {names}")


def remove_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for section in doc.sections:
        section.different_first_page_header_footer = False
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in list(footer.paragraphs):
                element = paragraph._element
                element.getparent().remove(element)


def set_columns(section, number: int, spacing_twips: int) -> None:
    sect_pr = section._sectPr
    columns = sect_pr.first_child_found_in("w:cols")
    if columns is None:
        columns = OxmlElement("w:cols")
        sect_pr.append(columns)
    columns.set(qn("w:num"), str(number))
    columns.set(qn("w:space"), str(spacing_twips))


def set_page_geometry(section, top: float) -> None:
    section.page_width = Inches(8.268)
    section.page_height = Inches(11.693)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def add_paragraph(doc: Document, text: str, paragraph_style, *, alignment=None, before=None, after=None):
    paragraph = doc.add_paragraph(style=paragraph_style)
    paragraph.add_run(text)
    if alignment is not None:
        paragraph.alignment = alignment
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    if after is not None:
        paragraph.paragraph_format.space_after = Pt(after)
    return paragraph


def add_heading(doc: Document, text: str, heading_style, level: int = 1):
    paragraph = add_paragraph(doc, text, heading_style, before=8 if level == 1 else 5, after=2)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        margin = tc_mar.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            tc_mar.append(margin)
        margin.set(qn("w:w"), "60")
        margin.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def write_cell(cell, value: str, text_style, *, bold: bool = False, center: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = text_style
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.size = Pt(7.5)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], widths: list[int], styles: dict):
    caption_paragraph = add_paragraph(doc, caption, styles["tablehead"], alignment=WD_ALIGN_PARAGRAPH.CENTER, before=5, after=2)
    caption_paragraph.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table, widths)
    header_cells = table.rows[0].cells
    for cell, text in zip(header_cells, headers):
        shade(cell, "D9E2F3")
        write_cell(cell, text, styles["tablecolhead"], bold=True, center=True)
    for record in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, record)):
            write_cell(cell, value, styles["tablecopy"], center=index > 0)
    set_table_width(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc: Document, path: Path, caption: str, number: int, styles: dict) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(3.12))
    cap = add_paragraph(doc, f"Fig. {number}. {caption}", styles["figurecaption"], alignment=WD_ALIGN_PARAGRAPH.CENTER, before=1, after=3)
    cap.paragraph_format.keep_together = True


def build() -> None:
    temporary_template = Path(tempfile.gettempdir()) / "ieee_a4_template_transitional.docx"
    strict_template_to_transitional(TEMPLATE, temporary_template)
    doc = Document(temporary_template)
    remove_template_body(doc)

    styles = {
        "title": style(doc, "paper title", "papertitle"),
        "author": style(doc, "Author", "author"),
        "abstract": style(doc, "Abstract", "abstract"),
        "keywords": style(doc, "Keywords", "keywords"),
        "heading1": style(doc, "Heading 1", "heading1"),
        "heading2": style(doc, "Heading 2", "heading2"),
        "body": style(doc, "Body Text", "bodytext"),
        "tablehead": style(doc, "table head", "tablehead"),
        "tablecolhead": style(doc, "table column head", "tablecolhead"),
        "tablecopy": style(doc, "table copy", "tablecopy"),
        "figurecaption": style(doc, "figure caption", "figurecaption"),
        "references": style(doc, "references", "References"),
    }

    title_section = doc.sections[0]
    set_page_geometry(title_section, 0.38)
    set_columns(title_section, 1, 720)

    title = add_paragraph(
        doc,
        "Forecasting IMDb Ratings from Movie Screenplays with Chunked SBERT and Stacked Regression",
        styles["title"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    title.paragraph_format.keep_with_next = True
    add_paragraph(doc, "Mahdin Mahboob", styles["author"], alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(doc, "Department of Computer Science and Engineering", styles["author"], alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(doc, "Southeast University, Dhaka, Bangladesh", styles["author"], alignment=WD_ALIGN_PARAGRAPH.CENTER, after=7)

    abstract_text = (
        "Abstract-Can a screenplay offer a useful early estimate of a film's later IMDb rating? "
        "We investigate that question on 5,195 feature-film scripts matched to rating and runtime metadata. "
        "The system represents each script with overlapping 256-word windows encoded by all-MiniLM-L6-v2, then combines the pooled 384-dimensional vector with 19 structural and metadata features. "
        "We compare mean prediction, metadata-only and structural baselines, TF-IDF with XGBoost, and weighted and unweighted SBERT-XGBoost variants. "
        "A Ridge stack is trained from out-of-fold predictions of four base models. In nested five-fold cross-validation, the stack reaches RMSE 0.935 +/- 0.032, MAE 0.706 +/- 0.028, and R2 0.577 +/- 0.010. "
        "It improves over its SBERT-XGBoost base learner by 0.013 MAE (paired Wilcoxon p = 4.17e-06). The metadata-only baseline still explains R2 = 0.382, so the result should not be read as screenplay quality in isolation. "
        "The experiment also shows that inverse-frequency weighting hurts this regression task: removing the weights lowers pooled MAE by 0.048 (p = 2.50e-26)."
    )
    add_paragraph(doc, abstract_text, styles["abstract"], after=4)
    add_paragraph(
        doc,
        "Index Terms-IMDb rating prediction, screenplay analysis, Sentence-BERT, stacked regression, XGBoost, SHAP.",
        styles["keywords"],
        after=4,
    )

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_page_geometry(body_section, 0.72)
    set_columns(body_section, 2, 360)

    add_heading(doc, "I. INTRODUCTION", styles["heading1"])
    for text in [
        "A screenplay is available long before audiences see the finished film, which makes it attractive for early-stage analysis. It is also an incomplete representation of the final product. Acting, direction, editing, publicity, release conditions, and the makeup of the eventual audience all affect ratings, but none is visible in a script. A useful model must therefore be evaluated as a limited forecasting tool rather than as an automated judge of artistic value.",
        "Script-based prediction has a clear precedent. Eliashberg, Hui, and Zhang used screenplay-derived textual features and kernel methods to estimate box-office performance at the green-light stage [1]. Our target is different: a continuous IMDb user rating. We focus on a practical modelling question: after basic metadata is accounted for, does screenplay text still improve a rating forecast?",
        "Long scripts create a technical obstacle. A standard sentence-transformer encoder cannot consume an entire feature screenplay in one pass. Sentence-BERT (SBERT) supplies semantic sentence representations that can be pooled efficiently [2], while long-context architectures such as Longformer address the same general problem through sparse attention [3]. Here, we use a simple chunk-and-pool design because it is inexpensive and already supported by the project code.",
        "The paper makes three modest contributions. First, it documents a reproducible SBERT plus structural-feature pipeline for a 5,195-script corpus. Second, it separates metadata, hand-crafted structure, bag-of-words text, and semantic embeddings through matched baselines. Third, it evaluates weighting and stacking with paired tests rather than treating a small point-estimate difference as conclusive evidence.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "II. RELATED WORK", styles["heading1"])
    for text in [
        "The early work in [1] is especially relevant because it uses information available before release. Later movie-prediction studies frequently add cast, budget, promotion, or online activity. Those signals may improve accuracy, but they answer a different question from a script-only analysis. The present work deliberately retains simple release-year, runtime, and decade variables because excluding them would make the evaluation less realistic, then measures their contribution explicitly.",
        "SBERT converts text segments into dense vectors suited to similarity and downstream prediction [2]. We use the 384-dimensional all-MiniLM-L6-v2 model, which keeps the embedding stage manageable for several thousand scripts. XGBoost [4] and LightGBM [5] are established tree-boosting implementations, while the remaining baselines are drawn from scikit-learn [9]. For interpretation, we use SHAP values [6] as diagnostic evidence, not as a causal account of how audiences rate films.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "III. DATASET AND TASK", styles["heading1"])
    for text in [
        "The project matches screenplay files obtained from the Internet Movie Script Database (IMSDb) [7] with a local metadata table containing IMDb ratings, release year, decade, and running time. The metadata sheet has 5,204 rows. Nine corresponding screenplay files are shorter than 1 KB and are excluded before modelling, leaving 5,195 usable records. The saved test-set manifest and experiment logs use this filtered count.",
        "The rating target ranges from 1.5 to 9.3, with mean 5.98, median 5.80, and standard deviation 1.44. Median runtime is 99 minutes. The distribution is concentrated in the middle of the scale, so error near very low or very high ratings is inherently less well supported. This imbalance motivates the weighting experiment in Section V, but it does not justify assuming that weighting will help a squared-error objective.",
        "The source is an archive rather than a random sample of released films. Older, well-known titles are more likely to survive in a screenplay archive, and release year is strongly associated with the target in the saved data. This is a data property, not evidence that time period causes higher audience ratings. The model must be interpreted within this archive population.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_table(
        doc,
        "TABLE I\nCORPUS AND FEATURE SUMMARY",
        ["Item", "Value"],
        [
            ["Metadata rows", "5,204"],
            ["Usable screenplays", "5,195"],
            ["Excluded scripts", "9 (<1 KB)"],
            ["Rating scale in corpus", "1.5 to 9.3"],
            ["Rating mean / median", "5.98 / 5.80"],
            ["Median runtime", "99 min"],
            ["Structured variables", "19"],
            ["SBERT representation", "384 dimensions"],
        ],
        [2500, 2100],
        styles,
    )

    add_heading(doc, "IV. METHOD", styles["heading1"])
    add_heading(doc, "A. Text Preparation and Features", styles["heading2"], level=2)
    for text in [
        "The implementation maintains two text views. Aggressive cleaning removes screenplay headers, stage directions, character cues, and most punctuation for TF-IDF and structural statistics. Light cleaning preserves case, punctuation, and sentence boundaries for SBERT. Keeping the two paths separate matters: the structural features benefit from normalization, whereas the pretrained encoder benefits from natural sentence form.",
        "Nineteen numeric variables are extracted from raw scripts. They include character, word, line, and sentence counts; vocabulary and long-word ratios; average sentence length and its variation; question and exclamation ratios; dialogue density; number of speaking characters; scene count; words per scene; and the three metadata fields. Missing numeric values are filled with a training-fold median and standardised using training-fold statistics only.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "B. Chunked SBERT Representation", styles["heading2"], level=2)
    for text in [
        "Each lightly cleaned script is split into 256-word windows with 50 words of overlap. The stride is therefore 206 words. Every window is encoded by all-MiniLM-L6-v2, and the document vector is the arithmetic mean of its window vectors. If a script produces K windows with embeddings e_k, its representation is e_bar = (1/K) sum(k=1 to K) e_k. This approximation preserves coverage of the document without truncating it to a single transformer context.",
        "The primary learner concatenates this 384-dimensional document vector with the scaled 19-variable feature vector. XGBoost is configured with histogram tree construction, learning rate 0.05, maximum depth 6, L1 regularisation 0.1, L2 regularisation 1.0, and early stopping after 20 rounds. Predictions are clipped to the valid 1-10 rating scale. These are implementation choices rather than a claim of globally optimal hyperparameters.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)
    add_figure(
        doc,
        FIGURES / "06_pipeline.png",
        "The implemented screenplay vectorization and prediction pipeline.",
        1,
        styles,
    )

    add_heading(doc, "C. Baselines, Weighting, and Stacking", styles["heading2"], level=2)
    for text in [
        "We compare six baselines: mean prediction, linear regression on metadata, linear regression on all structural variables, TF-IDF (8,000 one- and two-gram features) with XGBoost, weighted SBERT-XGBoost, and unweighted SBERT-XGBoost. The weighted variant assigns an inverse-frequency multiplier to four rating bands. In the stored 70/15/15 split, the rarest high-rating band receives a 13.24x multiplier; this is intentionally evaluated as an ablation rather than assumed to be beneficial.",
        "The stacked model uses four base predictions: metadata OLS, structural OLS, TF-IDF plus XGBoost, and unweighted SBERT plus XGBoost. A Ridge meta-regressor with alpha = 1.0 learns from inner-fold out-of-fold predictions. The outer five-fold loop then scores the stack on held-out records. This nesting prevents the meta-learner from seeing base predictions generated on the same target labels used to fit the base models.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)
    add_figure(
        doc,
        FIGURES / "07_stacking.png",
        "Nested stacking design used for the Ridge meta-regressor.",
        2,
        styles,
    )

    add_heading(doc, "D. Evaluation", styles["heading2"], level=2)
    for text in [
        "The initial comparison uses a fixed 70/15/15 split with random seed 42 (test n = 780). It reports RMSE, MAE, and R2 with 1,000 bootstrap resamples for 95% confidence intervals. Robustness is assessed with five shuffled folds. For paired comparisons, we apply a two-sided Wilcoxon signed-rank test to per-record absolute errors and report the bootstrap confidence interval for the MAE difference. This does not turn the data into an independent test set, but it does make the reported improvement more auditable.",
        "SHAP TreeExplainer is run on the fitted SBERT-XGBoost model. Because an individual embedding coordinate is not directly interpretable as a screenplay concept, SHAP plots are used to inspect model behaviour and identify possible dependence on metadata or individual vector dimensions. They are not used to infer causal screenplay properties.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "V. RESULTS", styles["heading1"])
    add_heading(doc, "A. Weighting and Main Baselines", styles["heading2"], level=2)
    add_paragraph(
        doc,
        "Table II reports the five-fold ablation. Metadata alone reaches R2 = 0.382, which establishes a substantial non-text benchmark. Adding the full structural set raises R2 to 0.440. TF-IDF with XGBoost performs similarly on R2 but is weaker than the semantic representation. The unweighted SBERT system is the best single model in this protocol, with RMSE 0.958 and MAE 0.720. In contrast, inverse-frequency weighting increases MAE from 0.720 to 0.768. The pooled difference is 0.048 MAE (95% CI: 0.039 to 0.057; p = 2.50e-26), so the degradation is unlikely to be a split artifact.",
        styles["body"],
        after=3,
    )
    add_table(
        doc,
        "TABLE II\nFIVE-FOLD BASELINE AND WEIGHTING ABLATION",
        ["Model", "RMSE", "MAE", "R2"],
        [
            ["Mean predictor", "1.438 +/- 0.056", "1.175 +/- 0.057", "0.000"],
            ["Metadata OLS", "1.130 +/- 0.047", "0.883 +/- 0.042", "0.382"],
            ["Structural OLS", "1.076 +/- 0.040", "0.833 +/- 0.031", "0.440"],
            ["TF-IDF + XGBoost", "1.083 +/- 0.034", "0.819 +/- 0.022", "0.433"],
            ["SBERT + XGB (weighted)", "1.000 +/- 0.027", "0.768 +/- 0.024", "0.516"],
            ["SBERT + XGB (unweighted)", "0.958 +/- 0.033", "0.720 +/- 0.027", "0.556"],
        ],
        [1900, 900, 900, 900],
        styles,
    )

    add_heading(doc, "B. Stacked Ensemble", styles["heading2"], level=2)
    add_paragraph(
        doc,
        "The separate nested stacking run is summarised in Table III. The strongest base learner, unweighted SBERT plus XGBoost, reaches RMSE 0.956 and MAE 0.719. The Ridge stack reduces these to 0.935 and 0.706, respectively, and obtains R2 = 0.577. The stack's pooled MAE advantage over its SBERT base model is 0.013 (95% CI: 0.008 to 0.019; p = 4.17e-06). The absolute gain is small, but its direction is stable across all five outer folds.",
        styles["body"],
        after=3,
    )
    add_table(
        doc,
        "TABLE III\nNESTED FIVE-FOLD STACKING RESULTS",
        ["Model", "RMSE", "MAE", "R2"],
        [
            ["Metadata OLS", "1.130 +/- 0.047", "0.883 +/- 0.042", "0.382"],
            ["Structural OLS", "1.076 +/- 0.040", "0.832 +/- 0.032", "0.440"],
            ["TF-IDF + XGBoost", "1.055 +/- 0.035", "0.790 +/- 0.028", "0.462"],
            ["SBERT + XGBoost", "0.956 +/- 0.034", "0.719 +/- 0.028", "0.558"],
            ["Ridge stack", "0.935 +/- 0.032", "0.706 +/- 0.028", "0.577"],
        ],
        [1900, 900, 900, 900],
        styles,
    )

    add_heading(doc, "C. Interpretation", styles["heading2"], level=2)
    for text in [
        "The saved SHAP report ranks several SBERT coordinates among the largest individual contributions in the XGBoost model. That is expected for a dense representation, but it does not give a plain-language explanation such as 'more dialogue improves rating.' At the feature level, the cleanest result is the baseline comparison: metadata alone has useful predictive power, and text improves it rather than replacing it.",
        "The main practical conclusion is therefore restrained. Script content contains information that helps forecast this archive's ratings, and combining complementary models produces a measurable increment. At the same time, the result is not high enough to support screening decisions without expert review, nor does it establish that a screenplay causes a specific audience response.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "VI. LIMITATIONS", styles["heading1"])
    for text in [
        "The corpus is archival and observational. It is subject to survivorship and selection effects, and its ratings are not evenly distributed across decades or score bands. The metadata table has 5,204 rows but nine files are filtered out by the 1 KB rule; future releases should publish a versioned manifest and hashes for the exact 5,195-file experiment set.",
        "The model sees no cast, director, budget, distribution, marketing, soundtrack, or finished-film qualities. It also pools chunks by an unweighted mean, which can obscure narrative order and long-range dependencies. Finally, a random cross-validation split allows closely related titles, eras, or archive patterns to appear in both training and test folds. Temporal and franchise-aware splits would be stronger tests of generalisation.",
        "The document reports artifacts that are present in the repository. It does not report the optional Word2Vec, GloVe-style, or MPNet paths as completed comparative evidence because the saved result tables do not contain a reproducible full benchmark for those alternatives. Those comparisons should be rerun and versioned before they are claimed in a future revision.",
    ]:
        add_paragraph(doc, text, styles["body"], after=3)

    add_heading(doc, "VII. CONCLUSION", styles["heading1"])
    add_paragraph(
        doc,
        "This study evaluates a practical method for forecasting IMDb ratings from screenplays. On 5,195 usable scripts, chunked SBERT vectors combined with structural variables outperform metadata, structural, and TF-IDF baselines. A nested Ridge stack produces the strongest result, with RMSE 0.935 and R2 0.577. The experiments also show why the result needs a careful reading: simple metadata explains a large fraction of the predictable variation, and inverse-frequency weighting makes performance worse. The next useful step is not a more ornate model description, but a versioned corpus and a stricter out-of-time evaluation.",
        styles["body"],
        after=4,
    )

    add_heading(doc, "REFERENCES", styles["heading1"])
    references = [
        "[1] J. Eliashberg, S. K. Hui, and Z. J. Zhang, \"Assessing box office performance using movie scripts: A kernel-based approach,\" IEEE Trans. Knowl. Data Eng., vol. 26, no. 11, pp. 2639-2648, 2014, doi: 10.1109/TKDE.2014.2306681.",
        "[2] N. Reimers and I. Gurevych, \"Sentence-BERT: Sentence embeddings using Siamese BERT-networks,\" in Proc. EMNLP-IJCNLP, 2019, pp. 3982-3992.",
        "[3] I. Beltagy, M. E. Peters, and A. Cohan, \"Longformer: The long-document transformer,\" arXiv:2004.05150, 2020.",
        "[4] T. Chen and C. Guestrin, \"XGBoost: A scalable tree boosting system,\" in Proc. 22nd ACM SIGKDD Int. Conf. Knowl. Discovery and Data Mining, 2016, pp. 785-794.",
        "[5] G. Ke et al., \"LightGBM: A highly efficient gradient boosting decision tree,\" in Proc. 31st Int. Conf. Neural Inf. Process. Syst., 2017, pp. 3149-3157.",
        "[6] S. M. Lundberg and S.-I. Lee, \"A unified approach to interpreting model predictions,\" in Adv. Neural Inf. Process. Syst., vol. 30, 2017.",
        "[7] Internet Movie Script Database, \"IMSDb,\" [Online]. Available: https://www.imsdb.com/. Accessed: Jul. 2026.",
        "[8] IMDb, \"IMDb datasets,\" [Online]. Available: https://developer.imdb.com/non-commercial-datasets/. Accessed: Jul. 2026.",
        "[9] F. Pedregosa et al., \"Scikit-learn: Machine learning in Python,\" J. Mach. Learn. Res., vol. 12, pp. 2825-2830, 2011.",
    ]
    for reference in references:
        add_paragraph(doc, reference, styles["references"], after=1)

    doc.core_properties.title = "Forecasting IMDb Ratings from Movie Screenplays"
    doc.core_properties.author = "Mahdin Mahboob"
    doc.core_properties.subject = "IEEE conference manuscript"
    doc.core_properties.comments = ""
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
