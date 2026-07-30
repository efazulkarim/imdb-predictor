"""Build wrapper that also supplies a deterministic table treatment."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_ieee_conference_paper as manuscript


def strict_template_to_transitional(source: Path, target: Path) -> None:
    replacements = {
        b"http://purl.oclc.org/ooxml/wordprocessingml/main": b"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        b"http://purl.oclc.org/ooxml/officeDocument/relationships": b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        b"http://purl.oclc.org/ooxml/drawingml/main": b"http://schemas.openxmlformats.org/drawingml/2006/main",
        b"application/vnd.ms-word.document.main+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
        b"application/vnd.ms-word.styles+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml",
        b"application/vnd.ms-word.fontTable+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml",
        b"application/vnd.ms-word.settings+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml",
        b"application/vnd.ms-word.numbering+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.numbering+xml",
        b"application/vnd.ms-word.footnotes+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        b"application/vnd.ms-word.endnotes+xml": b"application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
    }
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst:
        for member in src.infolist():
            data = src.read(member.filename)
            if member.filename.endswith((".xml", ".rels")):
                for old, new in replacements.items():
                    data = data.replace(old, new)
            dst.writestr(member, data)


def add_borders(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "808080")
        borders.append(border)
    table._tbl.tblPr.append(borders)


def add_table(doc, caption, headers, rows, widths, styles):
    cap = manuscript.add_paragraph(
        doc, caption, styles["tablehead"], alignment=WD_ALIGN_PARAGRAPH.CENTER, before=5, after=2
    )
    cap.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    add_borders(table)
    manuscript.set_table_width(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        manuscript.shade(cell, "D9E2F3")
        manuscript.write_cell(cell, text, styles["tablecolhead"], bold=True, center=True)
    for record in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, record)):
            manuscript.write_cell(cell, value, styles["tablecopy"], center=index > 0)
    manuscript.set_table_width(table, widths)
    doc.add_paragraph().paragraph_format.space_after = manuscript.Pt(1)
    return table


manuscript.strict_template_to_transitional = strict_template_to_transitional
manuscript.add_table = add_table
manuscript.build()
