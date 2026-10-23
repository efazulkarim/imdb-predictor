"""
Render paper/main.tex to Markdown and Word.

The .tex is the single source of truth; both outputs are generated from
it so the three versions cannot drift. Layout follows the IEEE
conference convention used by the reference submission: numbered
sections in Roman numerals, lettered subsections, "TABLE n." captions
above each table, "Fig. n." captions below each figure, and a two-column
body in the Word version with wide floats breaking out to full width.

Mathematics is kept as mathematics: the Markdown carries $...$ and
$$...$$ LaTeX, and the Word file carries native Office Math (OMML)
objects, produced by pandoc's LaTeX math reader. pandoc must be on PATH
or named by the PANDOC environment variable.

Requires python-docx and lxml. Run from the repository root:
    python build_paper_formats.py                         # paper/main.tex
    python build_paper_formats.py "paper/IEEE 6 pages.tex"
Outputs are written next to the .tex with the same stem.
"""

import copy
import io
import os
import re
import shutil
import subprocess
import sys
import zipfile

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

TEX = os.path.join('paper', 'main.tex')
OUT_MD = os.path.join('paper', 'main.md')
OUT_DOCX = os.path.join('paper', 'main.docx')

ROMAN = ['', 'I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X',
         'XI', 'XII']
LETTER = ' ABCDEFGHIJKLMNOPQRSTUVWXYZ'


# ------------------------------------------------------------------
# Inline LaTeX -> text
# ------------------------------------------------------------------
GREEK = {
    r'\lambda': 'lambda', r'\gamma': 'gamma', r'\varphi': 'phi',
    r'\pi': 'pi', r'\rho': 'rho', r'\tau': 'tau', r'\sigma': 'sigma',
    r'\Delta': 'Delta', r'\alpha': 'alpha', r'\mathcal': '',
    r'\varepsilon': 'eps', r'\bowtie': 'JOIN', r'\dagger': '+',
    r'\times': '×', r'\le': '<=', r'\ge': '>=', r'\in': 'in',
    r'\subseteq': '⊆', r'\rightarrow': '->', r'\to': '->',
    r'\approx': '≈', r'\pm': '±', r'\cdot': '·', r'\ldots': '...',
    r'\dots': '...', r'\exists': 'exists', r'\forall': 'for all',
    r'\neq': '≠', r'\equiv': '=', r'\notin': 'not in',
    r'\leftarrow': '<-', r'\hat': '', r'\bar': '',
}


def strip_math(s):
    """Turn inline math into readable plain text."""
    s = re.sub(r'\\mathrm\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\mathcal\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\text(?:bf|it|sc)?\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\emph\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\operatorname\{([^{}]*)\}', r'\1', s)
    # superscripts / subscripts
    s = re.sub(r'\^\{?2\}?', '^2', s)
    s = re.sub(r'\^\{([^{}]*)\}', r'^\1', s)
    s = re.sub(r'_\{([^{}]*)\}', r'_\1', s)
    s = re.sub(r'\\frac\{([^{}]*)\}\{([^{}]*)\}', r'(\1)/(\2)', s)
    # longest first, so \le cannot eat the front of \leftarrow
    for k in sorted(GREEK, key=len, reverse=True):
        s = s.replace(k, GREEK[k])
    s = s.replace('{,}', ',').replace('\\,', ' ').replace('\\;', ' ')
    s = s.replace('\\!', '').replace('\\ ', ' ')
    s = re.sub(r'[{}]', '', s)
    return s


ACCENTS = {
    r'{\o}': '\u00f8', r'\o': '\u00f8',
    r"{\'i}": '\u00ed', r"{\'a}": '\u00e1', r"{\'e}": '\u00e9',
    r"{\'o}": '\u00f3', r"{\'u}": '\u00fa', r"{\'y}": '\u00fd',
    r'{\"o}': '\u00f6', r'{\"a}': '\u00e4', r'{\"u}': '\u00fc',
    r'{\i}': '\u0131', r'{\ss}': '\u00df',
}


def plain_math(t):
    """Math rendered as ordinary text (numbers, or fallback contexts)."""
    t = t.replace('\\%', '%').replace('\\_', '_')
    return re.sub(r'\s+', ' ', strip_math(t)).strip()


def is_plain_number(t):
    """True for math that is only a number, range or percentage."""
    t = t.replace('{,}', ',').replace('\\%', '%').replace('\\,', ' ')
    return re.fullmatch(r'[\d\s.,%+\-\[\]()/:]*', t) is not None


def md_math(t):
    """LaTeX that is safe inside Markdown, including table cells."""
    t = t.replace('\\|', '\\Vert ').replace('|', '\\vert ')
    t = t.replace('\\textsc{', '\\mathrm{')
    return re.sub(r'\s+', ' ', t).strip()


# Inline math is lifted out of the text before any other substitution
# and replaced by \x01<n>\x02, so that text-level rewriting (dashes, %,
# brace stripping) cannot corrupt it. MATH holds the raw LaTeX.
MATH = []
MATH_TOKEN = re.compile('\x01(\\d+)\x02')


def inline(s, md=True, omml=False):
    """Convert an inline LaTeX fragment to Markdown or plain text.

    With omml=True, non-trivial math is left as MATH tokens for the DOCX
    writer to replace with Office Math objects.
    """
    s = re.sub(r'(?<!\\)%.*', '', s)          # comments

    def lift(m):
        MATH.append(m.group(1))
        return f'\x01{len(MATH) - 1}\x02'
    s = re.sub(r'(?<!\\)\$([^$]*)\$', lift, s)
    # structural macros that carry no body text
    s = re.sub(r'\\label\{[^}]*\}', '', s)
    s = re.sub(r'\\author\{(?:[^{}]|\{[^{}]*\})*\}', '', s)
    s = re.sub(r'\\(?:maketitle|IEEEoverridecommandlockouts|centering|'
               r'toprule|midrule|bottomrule|small|footnotesize|noindent)\b',
               '', s)
    for k, v in ACCENTS.items():
        s = s.replace(k, v)
    # "\ " is an escaped inter-word space; in the source it is often
    # followed by a line break rather than a literal space.
    s = re.sub(r'\\(\s)', r'\1', s)
    s = s.replace('\\,', '\u2009').replace('\\%', '%')
    s = s.replace('\\&', '&').replace('\\#', '#')
    s = s.replace('\\_', '_').replace('\\$', '$')

    # citations and cross-references
    s = re.sub(r'~?\\cite\{([^}]*)\}', lambda m: ' [' + ', '.join(
        str(BIBNUM.get(x.strip(), x.strip())) for x in m.group(1).split(',')) + ']', s)
    s = re.sub(r'~?\\(?:eq)?ref\{([^}]*)\}', lambda m: ' ' + _refname(m.group(1)), s)

    # font commands
    if md:
        s = re.sub(r'\\(?:hlt|textbf)\{([^{}]*)\}', r'**\1**', s)
        s = re.sub(r'\\emph\{([^{}]*)\}', r'*\1*', s)
        s = re.sub(r'\\textit\{([^{}]*)\}', r'*\1*', s)
        s = re.sub(r'\\texttt\{([^{}]*)\}', r'`\1`', s)
        s = re.sub(r'\\textsc\{([^{}]*)\}', r'\1', s)
    else:
        s = re.sub(r'\\(?:hlt|textbf|emph|textit|texttt|textsc)\{([^{}]*)\}',
                   r'\1', s)

    s = s.replace('---', '\u2014').replace('--', '\u2013')
    s = s.replace('~', ' ')
    s = s.replace("``", '\u201c').replace("''", '\u201d')
    s = re.sub(r'\\[a-zA-Z]+\b', '', s)       # leftover macros
    s = re.sub(r'[{}]', '', s)
    s = re.sub(r'\s+', ' ', s).strip()

    def put(m):
        t = MATH[int(m.group(1))]
        if is_plain_number(t):
            return plain_math(t)
        if md:
            return f'${md_math(t)}$'
        if omml:
            return m.group(0)
        return plain_math(t)
    return MATH_TOKEN.sub(put, s)


def eq_rows(env, src):
    """Split a display environment into its numbered rows of LaTeX."""
    src = re.sub(r'\\label\{[^}]*\}', '', src)
    if env.startswith('align'):
        rows = [r.replace('&', ' ') for r in re.split(r'\\\\', src)]
    else:
        rows = [src]
    return [re.sub(r'\s+', ' ', r).strip() for r in rows if r.strip()]


REFMAP = {}
BIBNUM = {}


def _refname(label):
    return REFMAP.get(label, label)


# ------------------------------------------------------------------
# Parse the .tex body into a block list
# ------------------------------------------------------------------
def parse(tex):
    body = tex.split(r'\begin{document}', 1)[1].split(r'\end{document}')[0]

    # bibitem keys -> reference numbers, needed before any \cite is rendered
    for n, key in enumerate(re.findall(r'\\bibitem\{([^}]*)\}', body), 1):
        BIBNUM[key] = n

    blocks = []
    sec_n, sub_n = 0, 0
    tab_n, fig_n, alg_n, eq_n = 0, 0, 0, 0

    # ---- first pass: assign numbers to every label ----
    for m in re.finditer(
            r'\\section\{|\\subsection\{|\\begin\{table\*?\}|'
            r'\\begin\{figure\*?\}|\\begin\{algorithm\}|'
            r'\\begin\{equation\}|\\begin\{align\}|\\label\{([^}]*)\}', body):
        tok = m.group(0)
        if tok.startswith(r'\section'):
            sec_n += 1
            sub_n = 0
            cur = ('sec', ROMAN[sec_n])
        elif tok.startswith(r'\subsection'):
            sub_n += 1
            cur = ('sub', f'{ROMAN[sec_n]}-{LETTER[sub_n]}')
        # Bare numbers: the prose already supplies "Table"/"Fig."/"Algorithm".
        elif tok.startswith(r'\begin{table'):
            tab_n += 1
            cur = ('tab', ROMAN[tab_n])
        elif tok.startswith(r'\begin{figure'):
            fig_n += 1
            cur = ('fig', str(fig_n))
        elif tok.startswith(r'\begin{algorithm'):
            alg_n += 1
            cur = ('alg', str(alg_n))
        elif tok.startswith(r'\begin{equation') or tok.startswith(r'\begin{align'):
            # LaTeX numbers every row of an align, so count rows.
            env = tok[len(r'\begin{'):-1]
            end = body.index(r'\end{%s}' % env, m.end())
            cur = ('eq', f'({eq_n + 1})')
            eq_n += len(eq_rows(env, body[m.end():end]))
        elif m.group(1):
            kind = m.group(1).split(':')[0]
            if kind in ('sec',):
                # the surrounding prose already supplies the word "Section"
                REFMAP[m.group(1)] = cur[1]
            elif kind == 'tab':
                REFMAP[m.group(1)] = cur[1]
            elif kind == 'fig':
                REFMAP[m.group(1)] = cur[1]
            elif kind == 'alg':
                REFMAP[m.group(1)] = cur[1]
            elif kind == 'eq':
                REFMAP[m.group(1)] = cur[1]

    # ---- second pass: emit blocks ----
    sec_n, sub_n, tab_n, fig_n, alg_n, eq_n = 0, 0, 0, 0, 0, 0
    i = 0
    while i < len(body):
        m = re.compile(
            r'\\title\{|\\begin\{abstract\}|\\begin\{IEEEkeywords\}|'
            r'\\section\*?\{|\\subsection\{|\\begin\{table\*?\}|'
            r'\\begin\{figure\*?\}|\\begin\{algorithm\}|'
            r'\\begin\{itemize\}|\\begin\{thebibliography\}|'
            r'\\begin\{(?:equation|align)\*?\}').search(body, i)
        if not m:
            blocks.append(('para', body[i:]))
            break
        if m.start() > i:
            blocks.append(('para', body[i:m.start()]))
        tok = m.group(0)

        if tok.startswith(r'\title'):
            txt, j = braced(body, m.end() - 1)
            blocks.append(('title', inline(txt)))
        elif tok.startswith(r'\begin{abstract}'):
            j = body.index(r'\end{abstract}', m.end())
            blocks.append(('abstract', body[m.end():j]))
            j += len(r'\end{abstract}')
        elif tok.startswith(r'\begin{IEEEkeywords}'):
            j = body.index(r'\end{IEEEkeywords}', m.end())
            blocks.append(('keywords', body[m.end():j]))
            j += len(r'\end{IEEEkeywords}')
        elif tok.startswith(r'\section'):
            txt, j = braced(body, m.end() - 1)
            if tok.endswith('*{'):
                blocks.append(('sec', ('', inline(txt))))
            else:
                sec_n += 1
                sub_n = 0
                blocks.append(('sec', (ROMAN[sec_n] + '.', inline(txt))))
        elif tok.startswith(r'\subsection'):
            txt, j = braced(body, m.end() - 1)
            sub_n += 1
            blocks.append(('sub', (LETTER[sub_n] + '.', inline(txt))))
        elif tok.startswith(r'\begin{table'):
            env = 'table*' if '*' in tok else 'table'
            j = body.index(r'\end{%s}' % env, m.end()) + len(r'\end{%s}' % env)
            tab_n += 1
            blocks.append(('table', (ROMAN[tab_n], body[m.end():j], '*' in tok)))
        elif tok.startswith(r'\begin{figure'):
            env = 'figure*' if '*' in tok else 'figure'
            j = body.index(r'\end{%s}' % env, m.end()) + len(r'\end{%s}' % env)
            fig_n += 1
            blocks.append(('figure', (fig_n, body[m.end():j], '*' in tok)))
        elif tok.startswith(r'\begin{algorithm}'):
            j = body.index(r'\end{algorithm}', m.end()) + len(r'\end{algorithm}')
            alg_n += 1
            blocks.append(('algorithm', (alg_n, body[m.end():j])))
        elif tok.startswith(r'\begin{itemize}'):
            j = body.index(r'\end{itemize}', m.end())
            blocks.append(('itemize', body[m.end():j]))
            j += len(r'\end{itemize}')
        elif tok.startswith(r'\begin{thebibliography}'):
            j = body.index(r'\end{thebibliography}', m.end())
            blocks.append(('bib', body[m.end():j]))
            j += len(r'\end{thebibliography}')
        else:  # equation / align
            env = re.match(r'\\begin\{([a-z]+\*?)\}', tok).group(1)
            j = body.index(r'\end{%s}' % env, m.end())
            rows = eq_rows(env, body[m.end():j])
            if env.endswith('*'):
                nums = [None] * len(rows)
            else:
                nums = list(range(eq_n + 1, eq_n + 1 + len(rows)))
                eq_n += len(rows)
            blocks.append(('equation', list(zip(rows, nums))))
            j += len(r'\end{%s}' % env)
        i = j
    return blocks


def braced(s, start):
    """Return (contents, index_after) for the brace group beginning at s[start]."""
    assert s[start] == '{'
    depth, i = 0, start
    while i < len(s):
        if s[i] == '{':
            depth += 1
        elif s[i] == '}':
            depth -= 1
            if depth == 0:
                return s[start + 1:i], i + 1
        i += 1
    raise ValueError('unbalanced braces')


def parse_tabular(src, md=True, omml=False, spans=False):
    """Extract (caption, rows) from a table environment body.

    md=False strips Markdown emphasis markers, which must not leak into
    DOCX cells where emphasis is carried by run formatting instead.
    spans=True makes each cell a (text, column span) pair, so that
    \\multicolumn headers and footnote rows can be merged in the DOCX.
    """
    cap = ''
    mc = re.search(r'\\caption\{', src)
    if mc:
        cap, _ = braced(src, mc.end() - 1)
    mt = re.search(r'\\begin\{tabular\}\{((?:[^{}]|\{[^{}]*\})*)\}', src)
    if not mt:
        return inline(cap, md, omml), []
    inner = src[mt.end():src.index(r'\end{tabular}', mt.end())]

    rows = []
    for raw in inner.split('\\\\'):
        raw = re.sub(r'\\(?:top|mid|bottom)rule', '', raw)
        raw = re.sub(r'\\cmidrule\(?[^)]*\)?\{[^}]*\}', '', raw)
        raw = re.sub(r'\\addlinespace(\[[^\]]*\])?', '', raw)
        if not raw.strip():
            continue
        cells = []
        for c in split_cells(raw):
            span = re.match(r'\s*\\multicolumn\{(\d+)\}', c)
            c = re.sub(r'\\multicolumn\{\d+\}\{[^{}]*\}\{(.*)\}', r'\1', c.strip(),
                       flags=re.S)
            c = re.sub(r'\\multirow\{[^{}]*\}\{[^{}]*\}\{(.*?)\}', r'\1', c, flags=re.S)
            c = c.replace(r'\footnotesize', '')
            t = inline(c, md, omml)
            cells.append((t, int(span.group(1)) if span else 1) if spans else t)
        if any(c[0] if spans else c for c in cells):
            rows.append(cells)
    return inline(cap, md, omml), rows


def split_cells(row):
    """Split a tabular row on & that are not inside braces."""
    out, depth, cur = [], 0, ''
    for ch in row:
        if ch == '{':
            depth += 1
        elif ch == '}':
            depth -= 1
        if ch == '&' and depth == 0:
            out.append(cur)
            cur = ''
        else:
            cur += ch
    out.append(cur)
    return out


# ------------------------------------------------------------------
# Markdown writer
# ------------------------------------------------------------------
def to_markdown(blocks):
    L = []
    for kind, payload in blocks:
        if kind == 'title':
            L += [f'# {payload}', '', '**Anonymous authors**', '']
        elif kind == 'abstract':
            L += [f'***Abstract**\u2014{inline(payload)}*', '']
        elif kind == 'keywords':
            L += [f'***Index Terms**\u2014{inline(payload)}*', '']
        elif kind == 'sec':
            num, txt = payload
            L += ['', f'## {num} {txt}'.strip(), '']
        elif kind == 'sub':
            num, txt = payload
            L += ['', f'### {num} {txt}', '']
        elif kind == 'para':
            for p in re.split(r'\n\s*\n', payload):
                t = inline(p)
                if t:
                    L += [t, '']
        elif kind == 'itemize':
            for item in payload.split(r'\item')[1:]:
                t = inline(item)
                if t:
                    L.append(f'- {t}')
            L.append('')
        elif kind == 'equation':
            for row, n in payload:
                tag = f' \\tag{{{n}}}' if n else ''
                L += ['$$', md_math(row) + tag, '$$', '']
        elif kind == 'table':
            num, src, _wide = payload
            cap, rows = parse_tabular(src)
            L += [f'**TABLE {num}.** {cap}', '']
            if rows:
                ncol = max(len(r) for r in rows)
                head = rows[0] + [''] * (ncol - len(rows[0]))
                L.append('| ' + ' | '.join(head) + ' |')
                L.append('|' + '---|' * ncol)
                for r in rows[1:]:
                    r = r + [''] * (ncol - len(r))
                    L.append('| ' + ' | '.join(r) + ' |')
            L.append('')
        elif kind == 'figure':
            num, src, _wide = payload
            cap, _ = parse_tabular(src)
            mi = re.search(r'\\includegraphics\[[^\]]*\]\{([^}]*)\}', src)
            path = mi.group(1).replace('.pdf', '.png') if mi else ''
            L += [f'![Fig. {num}](../{path})', '',
                  f'**Fig. {num}.** {cap}', '']
        elif kind == 'algorithm':
            num, src = payload
            mc = re.search(r'\\caption\{', src)
            cap = inline(braced(src, mc.end() - 1)[0]) if mc else ''
            L += [f'**Algorithm {num}.** {cap}', '']
            for n, ln in enumerate(re.findall(
                    r'\\(?:STATE|REQUIRE|ENSURE|REPEAT|UNTIL|RETURN)\b(.*)', src), 1):
                L.append(f'{n}: {inline(ln)}  ')
            L.append('')
        elif kind == 'bib':
            L += ['', '## References', '']
            items = payload.split(r'\bibitem')[1:]
            for n, it in enumerate(items, 1):
                _, rest = braced(it, it.index('{'))
                L.append(f'[{n}] {inline(it[it.index("}") + 1:])}')
            L.append('')
    return '\n'.join(L)


# ------------------------------------------------------------------
# DOCX writer
# ------------------------------------------------------------------
# Page geometry and type sizes copied from IEEEtran [conference,a4paper],
# so the Word file breaks pages where the LaTeX PDF does: 516pt text
# width in two 252pt columns, 672pt text height, 10pt type on 12pt lines.
PAGE_W, PAGE_H = 595.3, 841.9
MARGIN_LR, MARGIN_TOP = 39.65, 47.0
MARGIN_BOTTOM = PAGE_H - MARGIN_TOP - 672.0
COL_GAP = 12.0
TEXT_W = PAGE_W - 2 * MARGIN_LR
COL_W = (TEXT_W - COL_GAP) / 2
LEADING = {10: 12.0, 9: 10.5, 8: 9.5}


def set_columns(section, n):
    cols = section._sectPr.xpath('./w:cols')[0]
    cols.set(qn('w:num'), str(n))
    cols.set(qn('w:space'), str(int(COL_GAP * 20)))


def _border(parent, side, sz=6):
    b = OxmlElement(f'w:{side}')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), '0')
    b.set(qn('w:color'), '000000')
    parent.append(b)


def booktabs(table, header_rows):
    """Top and bottom rules on the table, a mid rule under the header."""
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    _border(borders, 'top', 8)
    _border(borders, 'bottom', 8)
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:w'), '60')                 # 3pt, close to \tabcolsep
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    # schema order: ... jc, tblBorders, shd, tblLayout, tblCellMar, tblLook
    layout = tblPr.find(qn('w:tblLayout'))
    look = tblPr.find(qn('w:tblLook'))
    (layout if layout is not None else look).addprevious(borders)
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)
    for cell in table.rows[header_rows - 1].cells:
        tcb = OxmlElement('w:tcBorders')
        _border(tcb, 'bottom', 4)
        cell._tc.get_or_add_tcPr().append(tcb)


def text_width(s, size):
    """Rough Times width of a cell, with math tokens at their plain length."""
    s = MATH_TOKEN.sub(lambda m: plain_math(MATH[int(m.group(1))]), s)
    return len(s) * size * 0.47


def column_widths(rows, ncol, avail, size=8):
    """Natural content widths, shrunk proportionally if they overflow."""
    nat = [12.0] * ncol
    for row in rows:
        c = 0
        for txt, span in row:
            if span == 1 and c < ncol:
                nat[c] = max(nat[c], text_width(txt, size) + 7)
            c += span
    tot = sum(nat)
    if tot > avail:
        nat = [w * avail / tot for w in nat]
    return nat


MNS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def find_pandoc():
    exe = os.environ.get('PANDOC') or shutil.which('pandoc')
    if not exe:
        sys.exit('pandoc not found: install it or set PANDOC=<path to pandoc>')
    return exe


def latex_to_omml(snippets):
    """Convert LaTeX math strings to <m:oMath> elements in one pandoc run."""
    if not snippets:
        return []
    # texmath rejects \Bigl/\Bigr (always paired in our sources, so
    # \left/\right is equivalent) and \textsc, and prefers \Vert over \|.
    def fix(s):
        s = s.replace('\\|', '\\Vert ').replace('\\textsc{', '\\mathrm{')
        s = s.replace('\\Bigl', '\\left').replace('\\Bigr', '\\right')
        return re.sub(r'\s+', ' ', s)
    src = '\n\n'.join('$$' + fix(s) + '$$' for s in snippets)
    out = subprocess.run([find_pandoc(), '-f', 'markdown', '-t', 'docx', '-o', '-'],
                         input=src.encode('utf-8'), capture_output=True,
                         check=True).stdout
    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read('word/document.xml'))
    maths = [p.find(f'{{{MNS}}}oMath') for p in root.iter(f'{{{MNS}}}oMathPara')]
    if len(maths) != len(snippets):
        sys.exit(f'pandoc returned {len(maths)} equations for {len(snippets)} inputs')
    return maths


def size_math(el, pt):
    """Give every math run an explicit font size (half-points)."""
    for r in el.iter(f'{{{MNS}}}r'):
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            # schema order inside m:r is m:rPr, w:rPr, then the text
            mrpr = r.find(f'{{{MNS}}}rPr')
            r.insert(0 if mrpr is None else r.index(mrpr) + 1, rpr)
        for old in rpr.findall(qn('w:sz')):
            rpr.remove(old)
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(pt * 2)))
        rpr.append(sz)


def to_docx(blocks, path):
    doc = Document()
    # (placeholder element, LaTeX, font size, display?) to be replaced
    # by Office Math once every snippet has been converted in one batch.
    pending = []

    def add_rich(p, text, size=10, bold=False, italic=False, font=None):
        """Add text to p, turning MATH tokens into math placeholders."""
        parts = MATH_TOKEN.split(text)
        for k, part in enumerate(parts):
            if k % 2:
                r = p.add_run()
                pending.append((r._r, MATH[int(part)], size, False))
                continue
            if not part:
                continue
            r = p.add_run(part)
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
            if font:
                r.font.name = font
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(10)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = Pt(LEADING[10])
    st.paragraph_format.first_line_indent = Pt(10)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # LaTeX hyphenates; without this Word sets visibly looser lines.
    hyph = OxmlElement('w:autoHyphenation')
    tab = doc.settings.element.find(qn('w:defaultTabStop'))
    tab.addnext(hyph)                       # schema position: after it

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(PAGE_W), Pt(PAGE_H)          # A4
    sec.left_margin = sec.right_margin = Pt(MARGIN_LR)
    sec.top_margin, sec.bottom_margin = Pt(MARGIN_TOP), Pt(MARGIN_BOTTOM)

    def para(text, size=10, bold=False, italic=False, align=None,
             indent=None, space_before=0, space_after=0, small_caps=False):
        p = doc.add_paragraph()
        add_rich(p, text, size, bold, italic)
        if small_caps:
            for r in p.runs:
                r.font.small_caps = True
        if align is not None:
            p.alignment = align
        pf = p.paragraph_format
        pf.first_line_indent = Pt(indent if indent is not None else 0)
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        if size in LEADING:
            pf.line_spacing = Pt(LEADING[size])
        else:
            pf.line_spacing = 1.0
        return p

    single_now = True   # title block starts single-column

    def ensure_columns(n):
        nonlocal single_now
        want_single = (n == 1)
        if want_single == single_now:
            return
        s = doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(s, n)
        single_now = want_single

    set_columns(sec, 1)

    for kind, payload in blocks:
        if kind == 'title':
            para(payload, size=22, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_after=8)
            para('Anonymous authors', size=11,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
        elif kind == 'abstract':
            ensure_columns(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Pt(10)
            p.paragraph_format.line_spacing = Pt(LEADING[9])
            r = p.add_run('Abstract\u2014')
            r.bold = True
            r.italic = True
            r.font.size = Pt(9)
            add_rich(p, inline(payload, md=False, omml=True), size=9, bold=True)
            p.paragraph_format.space_after = Pt(4)
        elif kind == 'keywords':
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(10)
            p.paragraph_format.line_spacing = Pt(LEADING[9])
            r = p.add_run('Index Terms\u2014')
            r.bold = True
            r.italic = True
            r.font.size = Pt(9)
            r2 = p.add_run(inline(payload, md=False))
            r2.bold = True
            r2.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(4)
        elif kind == 'sec':
            ensure_columns(2)
            num, txt = payload
            p = para(f'{num} {txt}'.strip(), size=10, small_caps=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8,
                     space_after=4)
            p.paragraph_format.keep_with_next = True
        elif kind == 'sub':
            num, txt = payload
            p = para(f'{num} {txt}', size=10, italic=True, space_before=4,
                     space_after=2)
            p.paragraph_format.keep_with_next = True
        elif kind == 'para':
            for chunk in re.split(r'\n\s*\n', payload):
                t = inline(chunk, md=False, omml=True)
                if t:
                    para(t, indent=10)
        elif kind == 'itemize':
            for item in payload.split(r'\item')[1:]:
                t = inline(item, md=False, omml=True)
                if t:
                    p = para('\u2022\t' + t, indent=-9)
                    p.paragraph_format.left_indent = Pt(18)
                    p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
        elif kind == 'equation':
            for row, n in payload:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf = p.paragraph_format
                pf.first_line_indent = Pt(0)
                pf.line_spacing = 1.0
                pf.space_before = Pt(3)
                pf.space_after = Pt(3)
                tex = row + (rf' \qquad ({n})' if n else '')
                pending.append((p._p, tex, 10, True))
        elif kind == 'table':
            num, src, wide = payload
            cap, rows = parse_tabular(src, md=False, omml=True, spans=True)
            ensure_columns(1 if wide else 2)
            p = para(f'Table {num}', size=8, small_caps=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6)
            p.paragraph_format.keep_with_next = True
            p = para(cap, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
            p.paragraph_format.keep_with_next = True
            if rows:
                ncol = max(sum(s for _, s in r) for r in rows)
                widths = column_widths(rows, ncol, TEXT_W if wide else COL_W)
                t = doc.add_table(rows=0, cols=ncol)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                t.autofit = False
                for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), widths):
                    gc.set(qn('w:w'), str(int(w * 20)))
                # the header is every row before the first data row whose
                # first cell is filled (group headers have an empty corner)
                header_rows = 1
                while (header_rows < len(rows) - 1
                       and not rows[header_rows - 1][0][0]):
                    header_rows += 1
                for ri, row in enumerate(rows):
                    cells = t.add_row().cells
                    c = 0
                    for txt, span in row:
                        if c >= ncol:
                            break
                        cell = cells[c]
                        if span > 1:
                            cell = cell.merge(cells[min(c + span, ncol) - 1])
                        elif len(row) == 1:                 # footnote row
                            cell = cell.merge(cells[ncol - 1])
                        cell.width = Pt(sum(widths[c:c + span]))
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.first_line_indent = Pt(0)
                        cp.paragraph_format.line_spacing = Pt(LEADING[8])
                        cp.alignment = (WD_ALIGN_PARAGRAPH.LEFT
                                        if c == 0 and span == 1 or len(row) == 1
                                        else WD_ALIGN_PARAGRAPH.CENTER)
                        add_rich(cp, txt, size=8, bold=(ri < header_rows),
                                 font='Times New Roman')
                        c += span
                booktabs(t, header_rows)
                para('', size=4, space_after=4).paragraph_format.line_spacing = Pt(4)
            if wide:
                ensure_columns(2)
        elif kind == 'figure':
            num, src, wide = payload
            cap, _ = parse_tabular(src, md=False, omml=True)
            mi = re.search(r'\\includegraphics\[[^\]]*\]\{([^}]*)\}', src)
            ensure_columns(1 if wide else 2)
            if mi:
                png = mi.group(1).replace('.pdf', '.png')
                if os.path.exists(png):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.add_run().add_picture(
                        png, width=Inches(6.8 if wide else 3.2))
            para(f'Fig. {num}.  {cap}', size=8,
                 align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
            if wide:
                ensure_columns(2)
        elif kind == 'algorithm':
            num, src = payload
            mc = re.search(r'\\caption\{', src)
            cap = inline(braced(src, mc.end() - 1)[0], md=False) if mc else ''
            para(f'Algorithm {num}.  {cap}', size=9, bold=True,
                 space_before=8, space_after=2)
            for n, ln in enumerate(re.findall(
                    r'\\(?:STATE|REQUIRE|ENSURE|REPEAT|UNTIL|RETURN)\b(.*)', src), 1):
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                add_rich(p, f'{n}:  {inline(ln, md=False, omml=True)}', size=9,
                         font='Consolas')
            para('', size=4, space_after=4)
        elif kind == 'bib':
            p = para('References', size=10, small_caps=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8,
                     space_after=4)
            p.paragraph_format.keep_with_next = True
            items = payload.split(r'\bibitem')[1:]
            for n, it in enumerate(items, 1):
                txt = inline(it[it.index('}') + 1:], md=False)
                p = para(f'[{n}]\t{txt}', size=8, space_after=1)
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.first_line_indent = Pt(-18)
                p.paragraph_format.tab_stops.add_tab_stop(Pt(18))

    maths = latex_to_omml([tex for _, tex, _, _ in pending])
    for (el, _, size, display), om in zip(pending, maths):
        om = copy.deepcopy(om)
        size_math(om, size)
        if display:
            mpara = etree.SubElement(el, f'{{{MNS}}}oMathPara')
            mpara.append(om)
        else:
            el.addprevious(om)
            el.getparent().remove(el)

    doc.save(path)


def main():
    global TEX, OUT_MD, OUT_DOCX
    if len(sys.argv) > 1:
        TEX = sys.argv[1]
        stem = os.path.splitext(TEX)[0]
        OUT_MD, OUT_DOCX = stem + '.md', stem + '.docx'
    tex = open(TEX, encoding='utf-8').read()
    blocks = parse(tex)
    print(f"parsed {len(blocks)} blocks; "
          f"{sum(1 for k, _ in blocks if k == 'table')} tables, "
          f"{sum(1 for k, _ in blocks if k == 'figure')} figures")

    md = to_markdown(blocks)
    open(OUT_MD, 'w', encoding='utf-8').write(md)
    print(f"wrote {OUT_MD} ({len(md.splitlines())} lines)")

    to_docx(blocks, OUT_DOCX)
    print(f"wrote {OUT_DOCX} ({os.path.getsize(OUT_DOCX) // 1024} KB)")
    print(f"{len(MATH)} inline math fragments, "
          f"{sum(len(p) for k, p in blocks if k == 'equation')} numbered display rows")


if __name__ == '__main__':
    main()
